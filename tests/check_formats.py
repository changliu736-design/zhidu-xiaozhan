# -*- coding: utf-8 -*-
"""实测多格式上传：.txt（GBK 编码）、.md、.docx，并验证出处跳转与正文预览接口。

会临时注册一个测试账号，测完把账号文件恢复原样、删掉测试目录。
"""
import json
import os
import shutil
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

BASE = "http://127.0.0.1:5000"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
USERS = os.path.join(ROOT, "backend", "data", "users.json")
BACKUP = USERS + ".fmt-bak"
UPLOADS = os.path.join(ROOT, "backend", "uploads")
WORK = os.path.join(ROOT, "tests", "_tmp_files")
PROBE_USER = "formatprobe"
PROBE_PASS = "probe123456"

results = []


def call(method, path, body=None, token=None, timeout=300):
    request = urllib.request.Request(
        BASE + path, data=json.dumps(body).encode() if body is not None else None, method=method)
    if body is not None:
        request.add_header("Content-Type", "application/json")
    if token:
        request.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.status, json.loads(response.read())
    except urllib.error.HTTPError as exc:
        raw = exc.read()
        try:
            return exc.code, json.loads(raw)
        except json.JSONDecodeError:
            return exc.code, {"error": raw.decode("utf-8", "replace")[:200]}
    except Exception as exc:  # noqa: BLE001
        return 0, {"error": f"{type(exc).__name__}: {exc}"}


def stream(path, token, timeout=300):
    request = urllib.request.Request(BASE + path)
    request.add_header("Authorization", "Bearer " + token)
    events, text = [], []
    with urllib.request.urlopen(request, timeout=timeout) as response:
        buffer = b""
        while True:
            block = response.read(1)
            if not block:
                break
            buffer += block
            if buffer.endswith(b"\n\n"):
                raw = buffer.decode("utf-8", "replace").strip()
                buffer = b""
                if raw.startswith("data:"):
                    event = json.loads(raw[5:].strip())
                    events.append(event)
                    if event.get("type") == "delta":
                        text.append(event["text"])
    return events, "".join(text)


def upload(token, filepath, filename=None):
    boundary = "----fmt" + uuid.uuid4().hex
    name = filename or os.path.basename(filepath)
    with open(filepath, "rb") as handle:
        content = handle.read()
    body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; '
            f'filename="{name}"\r\nContent-Type: application/octet-stream\r\n\r\n').encode() + content + \
           f"\r\n--{boundary}--\r\n".encode()
    request = urllib.request.Request(BASE + "/api/upload", data=body, method="POST")
    request.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
    request.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(request, timeout=300) as response:
            return response.status, json.loads(response.read())
    except urllib.error.HTTPError as exc:
        raw = exc.read()
        try:
            return exc.code, json.loads(raw)
        except json.JSONDecodeError:
            return exc.code, {"error": raw.decode("utf-8", "replace")[:200]}


def show(name, ok, detail=""):
    results.append((name, bool(ok), detail))
    print(f"  {'[通过]' if ok else '[失败]'} {name}" + (f"  — {detail}" if detail else ""))


PARAS = [
    ("研究背景", "医学图像分割是计算机辅助诊断中的基础任务。与自然图像不同，医学影像中病灶与周围组织的"
                 "对比度往往很低，边界划分因此格外困难。为了缓解这一问题，近年来的研究开始引入临床报告"
                 "作为额外的监督信号，希望借助语言中的先验知识提升分割精度。"),
    ("研究方法", "然而，把自由文本与像素级标注对齐仍然是一个开放问题。简单地把文本向量与视觉特征拼接，"
                 "常常导致网络干脆忽略掉语言模态，这一现象被称为模态坍塌。本文的做法是以语言表示作为条件，"
                 "重建被掩盖的视觉词元，从而迫使模型在跨模态之间保持语义一致。"),
    ("实验结果", "我们在三个公开数据集上评估了所提方法，包括胸部 X 光片与 CT 结肠成像数据。在所有设置下，"
                 "该框架都稳定优于强基线，Dice 相似系数平均提升 3.2 个百分点。消融实验进一步证实，"
                 "去掉文本分支或重建目标中的任意一个，性能都会明显下降。"),
    ("结论与展望", "最后我们讨论了本研究的局限性。该方法假设存在配对的报告文本，这一前提并非在所有临床"
                   "场景下都成立。未来工作将探索弱监督设置，并研究能否用大型语言模型生成足够忠实的伪报告，"
                   "用于训练分割网络。"),
]


def make_files():
    os.makedirs(WORK, exist_ok=True)
    paths = {}

    plain = "\n\n".join(f"{title}\n{body}" for title, body in PARAS)
    # 故意用 GBK 编码，测编码兜底
    txt_path = os.path.join(WORK, "分割研究备忘.txt")
    with open(txt_path, "wb") as handle:
        handle.write(plain.encode("gb18030"))
    paths["txt"] = txt_path

    md_lines = ["# 医学图像分割研究笔记", ""]
    for title, body in PARAS:
        md_lines += [f"## {title}", "", body, ""]
    md_lines += ["## 关键词", "", "跨模态、医学图像分割、条件重建", ""]
    md_path = os.path.join(WORK, "研究笔记.md")
    with open(md_path, "w", encoding="utf-8") as handle:
        handle.write("\n".join(md_lines))
    paths["md"] = md_path

    try:
        from docx import Document

        document = Document()
        document.add_heading("医学图像分割实验报告", level=0)
        for title, body in PARAS:
            document.add_heading(title, level=1)
            document.add_paragraph(body)
        document.add_heading("附：数据表", level=1)
        table = document.add_table(rows=2, cols=3)
        for i, value in enumerate(["数据集", "Dice", "提升"]):
            table.cell(0, i).text = value
        for i, value in enumerate(["QaTa-COV19", "0.842", "+3.1"]):
            table.cell(1, i).text = value
        docx_path = os.path.join(WORK, "实验报告.docx")
        document.save(docx_path)
        paths["docx"] = docx_path
    except ImportError:
        print("  （本机没有 python-docx，跳过 docx 测试）")

    # 老版 .doc 应当被拒绝
    doc_path = os.path.join(WORK, "旧版格式.doc")
    with open(doc_path, "wb") as handle:
        handle.write(b"\xd0\xcf\x11\xe0" + b"\x00" * 512)
    paths["doc"] = doc_path
    return paths


def wait_ready(token, doc_id, limit=240):
    deadline = time.time() + limit
    state = {}
    while time.time() < deadline:
        _, state = call("GET", f"/api/documents/{doc_id}/status", token=token)
        if state.get("index_status") in ("ready", "error"):
            return state
        time.sleep(3)
    return state


def main():
    shutil.copy2(USERS, BACKUP)
    print("已备份账号文件\n")
    status, data = call("POST", "/api/auth/register",
                        {"username": PROBE_USER, "password": PROBE_PASS})
    if status != 200:
        print("注册测试账号失败:", status, data)
        shutil.move(BACKUP, USERS)
        return 1
    token, user_id = data["token"], data["user"]["id"]
    user_dir = os.path.join(UPLOADS, user_id)
    print(f"测试账号就绪（{user_id}）\n")

    try:
        paths = make_files()
        print("=== 1. 生成测试文件 ===")
        for kind, path in paths.items():
            show(f"生成 {os.path.basename(path)}", os.path.exists(path),
                 f"{os.path.getsize(path)} 字节")

        print("\n=== 2. 各种格式上传 ===")
        uploaded = {}
        for kind in ("txt", "md", "docx"):
            if kind not in paths:
                continue
            status, data = upload(token, paths[kind])
            ok = status == 200 and data.get("success")
            doc = data.get("document", {})
            show(f"上传 {os.path.basename(paths[kind])}", ok,
                 f"{doc.get('page_count')} {doc.get('unit')} / {doc.get('chunk_count')} 片段 "
                 f"(source_type={doc.get('source_type')})" if ok else str(data)[:120])
            if ok:
                uploaded[kind] = doc
                show(f"  ↳ 单位是「{doc.get('unit')}」而不是「页」", doc.get("unit") == "段")

        status, data = upload(token, paths["doc"])
        show("老版 .doc 被拒绝", status == 400, f"HTTP {status} {str(data.get('error'))[:60]}")

        print("\n=== 3. 等索引建好 ===")
        for kind, doc in uploaded.items():
            state = wait_ready(token, doc["id"])
            show(f"{kind} 索引完成", state.get("index_status") == "ready",
                 json.dumps(state, ensure_ascii=False)[:120])

        if not uploaded:
            print("\n  三种格式都没上传成功，后续测试无法继续")
            return 1

        print("\n=== 4. 正文预览接口（点出处要用它）===")
        target = uploaded.get("docx") or uploaded.get("md") or uploaded.get("txt")
        status, blocks = call("GET", f"/api/documents/{target['id']}/blocks", token=token)
        show("接口返回成功", status == 200, f"HTTP {status}")
        show("带回单位名", blocks.get("unit") == "段", str(blocks.get("unit")))
        show("带 source_type", blocks.get("source_type") in ("text", "docx"), str(blocks.get("source_type")))
        show("有正文块", len(blocks.get("blocks", [])) > 0, f"{len(blocks.get('blocks', []))} 块")
        if blocks.get("blocks"):
            pages = sorted({b["page"] for b in blocks["blocks"]})
            show("块带连续编号", pages == list(range(1, len(pages) + 1)), f"编号 {pages}")
            show("块有正文", all(b.get("text") for b in blocks["blocks"]))
            if blocks.get("source_type") == "docx":
                joined = " ".join(b["text"] for b in blocks["blocks"])
                show("docx 标题被识别为章节", any(b.get("section") for b in blocks["blocks"]),
                     str([b.get("section") for b in blocks["blocks"] if b.get("section")][:2]))
                show("docx 表格文字也被提取", "QaTa-COV19" in joined)

        print("\n=== 5. 溯源问答（出处应当说「段」）===")
        question = "这篇文章用了什么方法？"
        events, answer = stream(
            f"/api/documents/{target['id']}/ask/stream?q=" + urllib.parse.quote(question), token)
        kinds = [e["type"] for e in events]
        done = [e for e in events if e["type"] == "done"]
        citations = done[0].get("citations", []) if done else []
        show("问答有流式输出", kinds.count("delta") > 0, f"{kinds.count('delta')} 个增量包")
        show("问答无错误", "error" not in kinds, str([e for e in events if e["type"] == "error"])[:120])
        show("答案带出处标记", "[[p." in answer, answer.strip()[:60])
        pages_in_range = all(1 <= c <= target["page_count"] for c in citations)
        show("出处编号在有效范围内", bool(citations) and pages_in_range,
             f"引用 {citations}，文档共 {target['page_count']} 段")
        print(f"      回答：{answer.strip()[:110]}…")

        print("\n=== 6. 摘要（也要带出处）===")
        events, summary = stream(f"/api/documents/{target['id']}/summary/stream", token)
        done = [e for e in events if e["type"] == "done"]
        show("摘要生成成功", len(summary.strip()) > 50, f"{len(summary)} 字")
        show("摘要带出处", bool(done and done[0].get("citations")),
             str(done[0].get("citations") if done else []))

        print("\n=== 7. PDF 仍照旧（回归检查）===")
        pdf = None
        for folder, _dirs, files in os.walk(UPLOADS):
            for name in files:
                if name.lower().endswith(".pdf") and "u_" in folder:
                    pdf = os.path.join(folder, name)
                    break
            if pdf:
                break
        if pdf:
            status, data = upload(token, pdf, "回归检查.pdf")
            show("PDF 仍可上传", status == 200 and data.get("success"),
                 f"单位={data.get('document', {}).get('unit')} "
                 f"类型={data.get('document', {}).get('source_type')}")
            if status == 200:
                show("PDF 单位仍是「页」", data["document"]["unit"] == "页")
                call("DELETE", f"/api/documents/{data['document']['id']}/delete", token=token)

    finally:
        print("\n=== 清理 ===")
        for kind, doc in list(locals().get("uploaded", {}).items()):
            call("DELETE", f"/api/documents/{doc['id']}/delete", token=token)
        shutil.move(BACKUP, USERS)
        if os.path.isdir(user_dir):
            shutil.rmtree(user_dir, ignore_errors=True)
        shutil.rmtree(WORK, ignore_errors=True)
        with open(USERS, encoding="utf-8") as handle:
            names = [v["username"] for v in json.load(handle).values()]
        show("账号文件已恢复原样", PROBE_USER not in names, f"账号 {names}")
        show("测试目录已删除", not os.path.exists(user_dir))
        leftovers = [n for n in os.listdir(UPLOADS) if n.startswith("u_") and n == user_id]
        show("uploads 里没留下测试目录", not leftovers)

    passed = sum(1 for _, ok, _ in results if ok)
    print(f"\n{'=' * 58}\n实测结果：{passed}/{len(results)} 项通过\n{'=' * 58}")
    for name, ok, detail in results:
        if not ok:
            print(f"  未通过：{name}  {detail}")
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
