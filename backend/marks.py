# -*- coding: utf-8 -*-
"""逻辑标记：记录"第 N 页第 M 段"，支持列表跳转与 Word 导出。

右侧是浏览器原生 PDF 阅读器，没法在上面叠加颜色，所以这里存的是
"位置 + 原文"，而不是图形层的红框。好处是跨 PDF/docx/txt 统一、可导出。
"""
import io
import json
import os
import re
import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from retrieval import extract_units


def marks_path(directory: Path, doc_id: str) -> Path:
    return directory / f"{doc_id}_marks.json"


def load_marks(directory: Path, doc_id: str) -> List[Dict]:
    path = marks_path(directory, doc_id)
    if not path.exists():
        return []
    try:
        with open(path, "r", encoding="utf-8") as handle:
            data = json.load(handle)
        return data if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError):
        return []


def save_marks(directory: Path, doc_id: str, marks: List[Dict]) -> None:
    path = marks_path(directory, doc_id)
    tmp = path.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8") as handle:
        json.dump(marks, handle, ensure_ascii=False, indent=2)
    tmp.replace(path)


def _is_content(text: str) -> bool:
    """过滤页码、装饰线和空白块：它们会被算成"第 1 段"，导致编号错位。"""
    stripped = text.strip()
    return len(stripped) >= 2 and not re.fullmatch(r"[-–—_·.\s\d]+", stripped)


def _int_to_rgb(value: int) -> Tuple[float, float, float]:
    return (((value >> 16) & 255) / 255.0,
            ((value >> 8) & 255) / 255.0,
            (value & 255) / 255.0)


def _pdf_block_infos(doc, page_no: int) -> List[Dict]:
    """返回某页的文本块：文字、矩形、字号、颜色。

    块序号与「第 N 段」的编号一致，标记、高亮、改字共用同一套定位。
    """
    import fitz

    if page_no < 1 or page_no > doc.page_count:
        raise ValueError(f"这份 PDF 只有 {doc.page_count} 页")
    page = doc[page_no - 1]
    raw_blocks = [b for b in page.get_text("blocks")
                  if len(b) >= 7 and b[6] == 0 and _is_content(str(b[4]))]

    spans = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                if span.get("text", "").strip():
                    spans.append((fitz.Rect(span["bbox"]),
                                  float(span.get("size", 10)),
                                  int(span.get("color", 0))))

    infos: List[Dict] = []
    for block in raw_blocks:
        text = str(block[4]).strip()
        rect = fitz.Rect(block[:4])
        matched = [(size, color) for inner, size, color in spans if inner.intersects(rect)]
        size = Counter(round(item[0], 1) for item in matched).most_common(1)[0][0] if matched else 10.0
        color_int = Counter(item[1] for item in matched).most_common(1)[0][0] if matched else 0
        infos.append({
            "text": text,
            "rect": rect,
            "size": size,
            "color": _int_to_rgb(color_int),
        })

    if infos:
        return infos

    # 极少见：PDF 没有可用的 block 结构，退回按空行切段并给出近似位置
    text = page.get_text("text").strip()
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if _is_content(p)]
    width = page.rect.width
    return [{"text": para,
             "rect": fitz.Rect(40, 60 + i * 18, width - 40, 78 + i * 18),
             "size": 10.0, "color": (0.0, 0.0, 0.0)}
            for i, para in enumerate(paras)]


def _pdf_paragraphs(source: Path, page_no: int) -> List[str]:
    import fitz

    with fitz.open(str(source)) as doc:
        return [block["text"] for block in _pdf_block_infos(doc, page_no)]


def _unit_paragraphs(source: Path, page_no: int) -> List[str]:
    units, _unit, _kind = extract_units(str(source))
    target = next((u for u in units if u.get("page") == page_no), None)
    if target is None:
        raise ValueError(f"找不到第 {page_no} 段")
    text = target.get("text", "")
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(paras) <= 1:
        paras = [p.strip() for p in text.split("\n") if p.strip()]
    return paras


def paragraph_text(source: Path, source_type: str, page_no: int, paragraph: int) -> str:
    if source_type == "pdf":
        paras = _pdf_paragraphs(source, page_no)
        where = f"第 {page_no} 页"
    else:
        paras = _unit_paragraphs(source, page_no)
        where = f"第 {page_no} 段"
    if not paras:
        raise ValueError(f"{where}没有可识别的段落")
    if paragraph < 1 or paragraph > len(paras):
        raise ValueError(f"{where}只有 {len(paras)} 个段落，没有第 {paragraph} 段")
    return paras[paragraph - 1]


def add_mark(directory: Path, info: Dict, source: Optional[Path],
             page: int, paragraph: int = 1, note: str = "") -> Dict:
    if source is None:
        raise ValueError("原始文件已丢失，无法读取原文段落")
    text = paragraph_text(source, info.get("source_type", "pdf"), page, paragraph)
    mark = {
        "id": uuid.uuid4().hex[:8],
        "doc_id": info.get("id", ""),
        "page": page,
        "paragraph": paragraph,
        "unit": info.get("unit", "页"),
        "text": text[:1200],
        "note": note.strip()[:500],
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }
    marks = load_marks(directory, info.get("id", ""))
    # 同一位置重复标记只保留一条，避免取消时删掉一条还剩一条
    for existing in marks:
        if existing.get("page") == page and existing.get("paragraph", 1) == paragraph:
            return existing
    marks.append(mark)
    marks.sort(key=lambda m: (m.get("page", 0), m.get("paragraph", 0)))
    save_marks(directory, info.get("id", ""), marks)
    return mark


def delete_mark(directory: Path, doc_id: str, mark_id: str) -> bool:
    marks = load_marks(directory, doc_id)
    remaining = [m for m in marks if m.get("id") != mark_id]
    if len(remaining) == len(marks):
        return False
    save_marks(directory, doc_id, remaining)
    return True


# ------------------------------------------------------------------ 导入
# 浏览器原生阅读器里的高亮只存在浏览器本地，服务端拿不到；
# 所以让用户把标注后的 PDF 导回来，解析注释转成标记。
_TEXT_ANNOT_CODES = {8, 9, 10, 11}  # highlight / underline / squiggly / strike-out


def _annot_quads(annot) -> List:
    import fitz

    quads = []
    try:
        vertices = annot.vertices or []
    except Exception:  # noqa: BLE001
        vertices = []
    for start in range(0, len(vertices) - 3, 4):
        try:
            quads.append(fitz.Quad(vertices[start:start + 4]).rect)
        except Exception:  # noqa: BLE001
            continue
    if not quads:
        quads.append(fitz.Rect(annot.rect))
    return quads


def _annot_block_indices(annot, blocks: List[Dict]) -> List[int]:
    """高亮覆盖了哪些文本块：每段线条归到重叠面积最大的块，跨段时返回多个。"""
    import fitz

    indices: List[int] = []
    for quad in _annot_quads(annot):
        winner = -1
        best_area = 0.0
        for index, block in enumerate(blocks):
            overlap = fitz.Rect(quad)
            overlap.intersect(fitz.Rect(block["rect"]))
            if overlap.is_empty:
                continue
            area = overlap.width * overlap.height
            if area > best_area:
                best_area = area
                winner = index
        if winner >= 0 and winner not in indices:
            indices.append(winner)
    return indices


def import_highlights(directory: Path, info: Dict, source: Path) -> Dict:
    """把 PDF 里手动画的高亮/下划线转成标记，纳入列表和 Word 导出。

    同一段落已有标记时不重复添加；一个高亮跨多段会拆成多条。
    """
    import fitz

    doc_id = info.get("id", "")
    try:
        doc = fitz.open(str(source))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"无法读取这份 PDF：{exc}")

    try:
        expected = int(info.get("page_count") or 0)
        if expected and doc.page_count != expected:
            raise ValueError(
                f"这份 PDF 有 {doc.page_count} 页，和原文献的 {expected} 页对不上，"
                f"请上传同一篇文献的标注版")

        marks = load_marks(directory, doc_id)
        seen = {(int(m.get("page", 0)), int(m.get("paragraph", 1))) for m in marks}
        additions: List[Dict] = []
        annotations = 0
        skipped = 0

        for page_index in range(doc.page_count):
            page = doc[page_index]
            annots = [a for a in (page.annots() or []) if a.type[0] in _TEXT_ANNOT_CODES]
            if not annots:
                continue
            blocks = _pdf_block_infos(doc, page_index + 1)
            for annot in annots:
                annotations += 1
                indices = _annot_block_indices(annot, blocks)
                if not indices:
                    skipped += 1
                    continue
                note = str((annot.info or {}).get("content") or "").strip()
                for index in indices:
                    key = (page_index + 1, index + 1)
                    if key in seen:
                        skipped += 1
                        continue
                    seen.add(key)
                    additions.append({
                        "id": uuid.uuid4().hex[:8],
                        "doc_id": doc_id,
                        "page": key[0],
                        "paragraph": key[1],
                        "unit": info.get("unit", "页"),
                        "text": blocks[index]["text"][:1200],
                        "note": note[:500],
                        "created_at": datetime.now().isoformat(timespec="seconds"),
                        "source": "manual",
                    })

        if additions:
            marks.extend(additions)
            marks.sort(key=lambda m: (m.get("page", 0), m.get("paragraph", 0)))
            save_marks(directory, doc_id, marks)
        return {"imported": len(additions), "skipped": skipped,
                "annotations": annotations, "marks": marks}
    finally:
        doc.close()


# ------------------------------------------------------------------ 编辑
def edits_path(directory: Path, doc_id: str) -> Path:
    return directory / f"{doc_id}_edits.json"


def edited_pdf_path(directory: Path, doc_id: str) -> Path:
    return directory / f"{doc_id}_edited.pdf"


def load_edits(directory: Path, doc_id: str) -> List[Dict]:
    path = edits_path(directory, doc_id)
    if not path.exists():
        return []
    try:
        with open(path, "r", encoding="utf-8") as handle:
            data = json.load(handle)
        return data if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError):
        return []


def save_edits(directory: Path, doc_id: str, edits: List[Dict]) -> None:
    path = edits_path(directory, doc_id)
    tmp = path.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8") as handle:
        json.dump(edits, handle, ensure_ascii=False, indent=2)
    tmp.replace(path)


def version_of(directory: Path, doc_id: str) -> str:
    """标记/编辑的最后修改时间，用作前端刷新 PDF 的缓存版本号。"""
    stamps = [path.stat().st_mtime_ns
              for path in (marks_path(directory, doc_id), edits_path(directory, doc_id))
              if path.exists()]
    return str(max(stamps)) if stamps else ""


_COLOR_NAMES = {
    "红": (1.0, 0.0, 0.0),
    "蓝": (0.0, 0.15, 0.9), "绿": (0.0, 0.55, 0.15),
    "黑": (0.0, 0.0, 0.0), "橙": (1.0, 0.5, 0.0),
    "紫": (0.55, 0.0, 0.75), "灰": (0.4, 0.4, 0.4),
}

_FONT_CHOICES = {
    "宋体": ("china-s", None),
    "黑体": ("china-ss", None),
    "微软雅黑": ("msyh", r"C:\Windows\Fonts\msyh.ttc"),
    "雅黑": ("msyh", r"C:\Windows\Fonts\msyh.ttc"),
    "楷体": ("kai", r"C:\Windows\Fonts\simkai.ttf"),
}


def _parse_color(value) -> Optional[Tuple[float, float, float]]:
    if isinstance(value, (list, tuple)) and len(value) == 3:
        try:
            return tuple(max(0.0, min(1.0, float(item))) for item in value)  # type: ignore[return-value]
        except (TypeError, ValueError):
            return None
    if not isinstance(value, str):
        return None
    text = value.strip().lower()
    if re.fullmatch(r"#[0-9a-f]{6}", text):
        return tuple(int(text[i:i + 2], 16) / 255.0 for i in (1, 3, 5))  # type: ignore[return-value]
    for name, color in _COLOR_NAMES.items():
        if name in text:
            return color
    return None


def _resolve_font(value) -> Tuple[str, Optional[str]]:
    text = str(value or "")
    for name, spec in _FONT_CHOICES.items():
        if name in text:
            return spec
    return "china-s", None


def _block_for(doc, page_no: int, paragraph: int) -> Dict:
    if page_no < 1 or page_no > doc.page_count:
        raise ValueError(f"这份 PDF 只有 {doc.page_count} 页")
    blocks = _pdf_block_infos(doc, page_no)
    if paragraph < 1 or paragraph > len(blocks):
        raise ValueError(f"第 {page_no} 页只有 {len(blocks)} 个段落，没有第 {paragraph} 段")
    return blocks[paragraph - 1]


def add_edit(directory: Path, info: Dict, source: Optional[Path],
             payload: Dict) -> Dict:
    if source is None:
        raise ValueError("原始文件已丢失，无法修改")
    if info.get("source_type", "pdf") != "pdf":
        raise ValueError("目前只支持在 PDF 上改字/改色")
    try:
        page = int(payload.get("page") or 0)
        paragraph = int(payload.get("paragraph") or 1)
    except (TypeError, ValueError):
        raise ValueError("页码或段落编号不是数字")

    kind = payload.get("type")
    if kind not in ("replace", "recolor"):
        raise ValueError("不支持的修改类型")

    import fitz

    with fitz.open(str(source)) as doc:
        block = _block_for(doc, page, paragraph)

    if kind == "replace":
        text = str(payload.get("text") or "").strip()
        if not text:
            raise ValueError("没有收到要写入的新文字")
    else:
        text = block["text"]

    color = _parse_color(payload.get("color")) or block["color"]
    size = payload.get("size")
    try:
        size = float(size) if size else None
    except (TypeError, ValueError):
        size = None

    edit = {
        "id": uuid.uuid4().hex[:8],
        "type": kind,
        "page": page,
        "paragraph": paragraph,
        "text": text[:4000],
        "original": block["text"][:4000],
        "color": list(color),
        "font": str(payload.get("font") or "") or None,
        "size": size,
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }

    edits = [item for item in load_edits(directory, info.get("id", ""))
             if not (item.get("page") == page and item.get("paragraph") == paragraph)]
    edits.append(edit)
    edits.sort(key=lambda item: (item.get("page", 0), item.get("paragraph", 0)))
    save_edits(directory, info.get("id", ""), edits)
    return edit


def delete_edit(directory: Path, doc_id: str, edit_id: str) -> bool:
    edits = load_edits(directory, doc_id)
    remaining = [item for item in edits if item.get("id") != edit_id]
    if len(remaining) == len(edits):
        return False
    save_edits(directory, doc_id, remaining)
    return True


def _build_edited_pdf(source: Path, out: Path, tmp: Path,
                      marks: List[Dict], edits: List[Dict]) -> None:
    import fitz

    doc = fitz.open(str(source))
    cache: Dict[Tuple[int, int], Dict] = {}

    def block_of(page: int, paragraph: int) -> Dict:
        key = (page, paragraph)
        if key not in cache:
            cache[key] = _block_for(doc, page, paragraph)
        return cache[key]

    # 先把所有定位算完，再动页面：redact 会改变页面上的文本结构
    for item in list(edits) + list(marks):
        block_of(int(item["page"]), int(item["paragraph"]))

    try:
        for edit in edits:
            page_no = int(edit["page"])
            block = cache[(page_no, int(edit["paragraph"]))]
            page = doc[page_no - 1]
            rect = fitz.Rect(block["rect"])
            text = edit.get("text") or block["text"]
            color = tuple(edit.get("color") or block["color"])
            size = float(edit.get("size") or block["size"] or 10.0)
            fontname, fontfile = _resolve_font(edit.get("font"))

            page.add_redact_annot(rect, fill=(1, 1, 1))
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)

            box = fitz.Rect(rect.x0, rect.y0 - 0.5, rect.x1 + 2, rect.y1 + 0.5)
            kwargs = {"fontname": fontname, "color": color,
                      "align": fitz.TEXT_ALIGN_LEFT}
            if fontfile:
                kwargs["fontfile"] = fontfile

            current = size
            placed = False
            while current >= 4.5:
                kwargs["fontsize"] = current
                if page.insert_textbox(box, text, **kwargs) >= 0:
                    placed = True
                    break
                current *= 0.92
            if not placed:
                # 最小字号还放不下：向下扩一点，宁可轻微挤压也不丢内容
                wide = fitz.Rect(box.x0, box.y0, box.x1,
                                 min(page.rect.y1 - 2, box.y1 + rect.height))
                kwargs["fontsize"] = 4.5
                page.insert_textbox(wide, text, **kwargs)

        for mark in marks:
            page_no = int(mark["page"])
            block = cache[(page_no, int(mark["paragraph"]))]
            page = doc[page_no - 1]
            annot = page.add_highlight_annot(fitz.Rect(block["rect"]))
            annot.set_colors(stroke=(1, 0, 0))
            annot.set_opacity(0.45)
            annot.update()

        # 外挂中文字体会整份嵌入（几 MB），按实际用到的字形做子集后通常只剩一百多 KB
        try:
            doc.subset_fonts()
        except Exception:  # noqa: BLE001
            pass
        doc.save(str(tmp), garbage=3, deflate=True)
    finally:
        doc.close()
    os.replace(tmp, out)


def ensure_edited_pdf(directory: Path, doc_id: str, source: Optional[Path]) -> Optional[Path]:
    """按需生成带红色高亮/改字结果的 PDF 副本；没有修改时返回 None。"""
    if source is None:
        return None
    marks = load_marks(directory, doc_id)
    edits = load_edits(directory, doc_id)
    if not marks and not edits:
        return None

    out = edited_pdf_path(directory, doc_id)
    tmp = directory / f"{doc_id}_edited.tmp.pdf"
    stamps = [path.stat().st_mtime_ns
              for path in (marks_path(directory, doc_id), edits_path(directory, doc_id))
              if path.exists()]
    newest = max(stamps) if stamps else 0
    if out.exists() and out.stat().st_mtime_ns >= newest:
        return out
    try:
        _build_edited_pdf(source, out, tmp, marks, edits)
        return out
    except Exception as exc:  # noqa: BLE001  生成失败时退回原文件，至少能继续阅读
        print(f"[编辑] 生成修改版 PDF 失败：{type(exc).__name__}: {exc}")
        return out if out.exists() else None


def build_docx(info: Dict, marks: List[Dict]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    try:
        style = document.styles["Normal"]
        style.font.name = "宋体"
        style.font.size = Pt(10.5)
        style.element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "宋体")
    except Exception:  # noqa: BLE001  字体设置失败不影响正文导出
        pass

    title = info.get("display_name") or info.get("original_filename", "文献")
    document.add_heading(f"《{title}》阅读标记", level=1)
    document.add_paragraph(
        f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}　共 {len(marks)} 条标记")

    for index, mark in enumerate(marks, 1):
        unit = mark.get("unit", info.get("unit", "页"))
        origin = "（手动高亮）" if mark.get("source") == "manual" else ""
        document.add_heading(
            f"标记 {index}　第 {mark.get('page', '?')} {unit} 第 {mark.get('paragraph', 1)} 段{origin}",
            level=2)
        quote = document.add_paragraph(mark.get("text", ""))
        quote.paragraph_format.left_indent = Pt(12)
        if mark.get("note"):
            document.add_paragraph(f"备注：{mark['note']}")
        document.add_paragraph(f"记录于 {mark.get('created_at', '')}")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
